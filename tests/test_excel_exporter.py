import os
import tempfile
from datetime import date
import openpyxl
import pytest

from app.utils.excel_exporter import (
    export_invoices_to_excel,
    export_gstr1_to_excel,
    export_areas_to_excel,
    export_receivables_to_excel,
    export_payments_to_excel,
    export_customer_directory_to_excel,
    export_items_to_excel,
)


class MockCustomer:
    name = "Shubham"
    mobile = "9876543210"


class MockPayment:
    amount = 150000.0


class MockInvoice:
    invoice_number = "INV-2026-0001"
    customer = MockCustomer()
    invoice_date = date(2026, 9, 9)
    due_date = date(2026, 9, 20)
    status = "PARTIALLY PAID"
    grand_total = 200000.0
    payments = [MockPayment()]


def test_export_invoices_excel():
    with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as td:
        xlsx_path = os.path.join(td, "test_invoices.xlsx")
        export_invoices_to_excel(xlsx_path, [MockInvoice()])

        assert os.path.exists(xlsx_path)
        assert os.path.getsize(xlsx_path) > 1000

        wb = openpyxl.load_workbook(xlsx_path)
        try:
            ws = wb.active
            assert "INVOICE SALES" in ws["A1"].value
            assert len(ws._charts) >= 1
        finally:
            wb.close()


def test_export_invoices_csv_fallback():
    with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as td:
        csv_path = os.path.join(td, "test_invoices.csv")
        export_invoices_to_excel(csv_path, [MockInvoice()])
        assert os.path.exists(csv_path)
        with open(csv_path, "r", encoding="utf-8-sig") as f:
            content = f.read()
            assert "INV-2026-0001" in content
            assert "Shubham" in content


def test_export_areas_excel_with_chart():
    with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as td:
        xlsx_path = os.path.join(td, "test_areas.xlsx")
        areas = [
            {"area": "BEDROOM", "item_count": 5, "revenue": 120000.0, "percent": 60.0},
            {"area": "LIVING ROOM", "item_count": 3, "revenue": 80000.0, "percent": 40.0},
        ]
        export_areas_to_excel(xlsx_path, areas)
        assert os.path.exists(xlsx_path)
        wb = openpyxl.load_workbook(xlsx_path)
        try:
            ws = wb.active
            assert "FURNITURE SALES BY ROOM AREA" in ws["A1"].value
            assert len(ws._charts) >= 1
        finally:
            wb.close()


def test_export_receivables_excel_with_chart():
    with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as td:
        xlsx_path = os.path.join(td, "test_receivables.xlsx")
        custs = [
            {
                "name": "Shubham", "mobile": "9876543210", "gstin": "-", "city": "Thane",
                "invoices_count": 1, "total_billed": 200000.0, "total_paid": 150000.0,
                "balance_due": 50000.0, "compliance_pct": 75.0
            }
        ]
        export_receivables_to_excel(xlsx_path, custs)
        assert os.path.exists(xlsx_path)
        wb = openpyxl.load_workbook(xlsx_path)
        try:
            ws = wb.active
            assert len(ws._charts) >= 1
        finally:
            wb.close()


def test_export_gstr1_excel():
    with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as td:
        xlsx_path = os.path.join(td, "test_gstr1.xlsx")
        rows = [
            {
                "invoice_number": "INV-001", "date": date(2026, 9, 9),
                "customer_name": "Shubham", "customer_mobile": "9876543210",
                "customer_gstin": "27AABCS1429B1Z", "customer_state": "Maharashtra",
                "is_b2b": True, "taxable_value": 200000.0, "gst_rate": 18.0,
                "cgst": 18000.0, "sgst": 18000.0, "igst": 0.0,
                "grand_total": 236000.0, "paid_amount": 150000.0,
                "balance": 86000.0, "status": "PARTIALLY PAID"
            }
        ]
        export_gstr1_to_excel(xlsx_path, rows)
        assert os.path.exists(xlsx_path)
        wb = openpyxl.load_workbook(xlsx_path)
        try:
            ws = wb.active
            assert "GSTR-1 TAX REGISTER" in ws["A1"].value
        finally:
            wb.close()


def test_export_payments_and_customer_directory():
    with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as td:
        p_path = os.path.join(td, "test_payments.xlsx")
        pays = [{
            "date": date(2026, 9, 9), "invoice_number": "INV-001",
            "customer_name": "Shubham", "customer_mobile": "9876543210",
            "mode": "UPI", "reference": "ADV-SHUBHAM-150K", "amount": 150000.0, "notes": "Advance"
        }]
        export_payments_to_excel(p_path, pays)
        assert os.path.exists(p_path)

        c_path = os.path.join(td, "test_cust.xlsx")
        custs = [{
            "id": 16, "name": "Shubham", "mobile": "9876543210", "email": "s@example.com",
            "city": "Thane", "state": "MH", "gstin": "-", "invoice_count": 1,
            "total_invoiced": 200000.0, "total_paid": 150000.0, "outstanding": 50000.0,
            "is_settled": False
        }]
        export_customer_directory_to_excel(c_path, custs)
        assert os.path.exists(c_path)


def test_export_items_excel():
    with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as td:
        i_path = os.path.join(td, "test_items.xlsx")
        items = [
            {"area": "BEDROOM", "description": "Wardrobe", "size": "10x9", "qty_raw": "1", "rate_raw": "120000", "amount": 120000.0},
            {"area": "LIVING ROOM", "description": "TV Unit", "size": "8x7", "qty_raw": "1", "rate_raw": "80000", "amount": 80000.0}
        ]
        export_items_to_excel(i_path, items)
        assert os.path.exists(i_path)
